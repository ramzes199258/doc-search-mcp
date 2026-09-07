# -*- coding: utf-8 -*-
"""MCP-сервер поиска по документам: инкремент + гибрид + OCR + таблицы."""
import os
import re
import json
import math
import hashlib
import threading
import time
from pathlib import Path
from collections import Counter

import numpy as np
import faiss
import fitz
import requests
from docx import Document as DocxDocument
from watchdog.observers import Observer
from watchdog.events import FileSystemEventHandler

# ---------------- НАСТРОЙКИ (из переменных окружения) ----------------
DOCS_DIR = Path(os.environ.get("DOCS_DIR", r"C:\doc-mcp\docs"))
INDEX_DIR = Path(os.environ.get("INDEX_DIR", r"C:\doc-mcp\index"))
CACHE_DIR = INDEX_DIR / "cache"
OLLAMA_URL = os.environ.get("OLLAMA_URL", "http://127.0.0.1:11434")
EMBED_MODEL = os.environ.get("EMBED_MODEL", "emb-cpu")
CHUNK_SIZE = int(os.environ.get("CHUNK_SIZE", "1000"))
CHUNK_OVERLAP = int(os.environ.get("CHUNK_OVERLAP", "150"))
PORT = int(os.environ.get("PORT", "8000"))
HOST = os.environ.get("HOST", "0.0.0.0")
REINDEX_DELAY = int(os.environ.get("REINDEX_DELAY", "5"))
EMBED_BATCH = int(os.environ.get("EMBED_BATCH", "16"))
OCR_ENABLED = os.environ.get("OCR_ENABLED", "true").lower() == "true"
OCR_DPI = int(os.environ.get("OCR_DPI", "200"))
OCR_MIN_TEXT = int(os.environ.get("OCR_MIN_TEXT", "40"))
# ---------------------------------------------------------------------

try:
    from mcp.server.fastmcp import FastMCP
except ImportError:
    from mcp.server import FastMCP

try:
    mcp = FastMCP("doc-search", host=HOST, port=PORT)
except TypeError:
    os.environ["FASTMCP_HOST"] = HOST
    os.environ["FASTMCP_PORT"] = str(PORT)
    mcp = FastMCP("doc-search")

INDEX = None
META = []
BM25_OBJ = None
LAST_STATE = {}
OCR_ENGINE = None
_reindex_lock = threading.Lock()
SUPPORTED = (".txt", ".md", ".pdf", ".docx")


def tokenize(s: str):
    return re.findall(r"[a-zа-яё0-9]+", s.lower())


class BM25:
    def __init__(self, docs):
        self.N = len(docs)
        self.tf = []
        self.df = Counter()
        self.dl = []
        for d in docs:
            t = Counter(tokenize(d))
            self.tf.append(t)
            self.dl.append(sum(t.values()))
            self.df.update(t.keys())
        total = sum(self.dl)
        self.avgdl = (total / self.N) if self.N else 1.0

    def score(self, query):
        s = [0.0] * self.N
        k1, b = 1.5, 0.75
        for tok in tokenize(query):
            df = self.df.get(tok, 0)
            if not df:
                continue
            idf = math.log(1 + (self.N - df + 0.5) / (df + 0.5))
            for i in range(self.N):
                f = self.tf[i].get(tok, 0)
                if f:
                    s[i] += idf * (f * (k1 + 1)) / (f + k1 * (1 - b + b * self.dl[i] / self.avgdl))
        return s


def build_bm25():
    global BM25_OBJ
    try:
        BM25_OBJ = BM25([m["text"] for m in META]) if META else None
    except Exception as e:
        print("BM25 не собрался, остаёмся на векторном поиске:", repr(e), flush=True)
        BM25_OBJ = None


def get_ocr():
    global OCR_ENGINE
    if OCR_ENGINE is None:
        from rapidocr_onnxruntime import RapidOCR
        OCR_ENGINE = RapidOCR()
    return OCR_ENGINE


def ocr_page(page):
    pix = page.get_pixmap(dpi=OCR_DPI)
    img = np.frombuffer(pix.samples, dtype=np.uint8).reshape(pix.height, pix.width, pix.n)
    if pix.n == 4:
        img = img[:, :, :3]
    result, _ = get_ocr()(img)
    if not result:
        return ""
    return " ".join(line[1] for line in result)


def extract_text(path: Path) -> str:
    suf = path.suffix.lower()
    if suf in (".txt", ".md"):
        return path.read_text(encoding="utf-8", errors="ignore")
    if suf == ".pdf":
        parts = []
        with fitz.open(str(path)) as doc:
            for i, page in enumerate(doc):
                text = page.get_text()
                try:
                    for t in page.find_tables().tables:
                        for row in t.extract():
                            parts.append(" | ".join(c or "" for c in row))
                except Exception:
                    pass
                if OCR_ENABLED and len(text.strip()) < OCR_MIN_TEXT:
                    print(f"  OCR: страница {i + 1}...", flush=True)
                    text = ocr_page(page)
                parts.append(text)
        return "\n".join(parts)
    if suf == ".docx":
        doc = DocxDocument(str(path))
        return "\n".join(p.text for p in doc.paragraphs)
    return ""


def chunk_text(text: str):
    text = re.sub(r"\s+", " ", text).strip()
    chunks, start = [], 0
    while start < len(text):
        chunks.append(text[start:start + CHUNK_SIZE])
        start += CHUNK_SIZE - CHUNK_OVERLAP
    return [c for c in chunks if len(c) > 50]


def embed(texts, retries=3):
    import time
    for attempt in range(retries):
        try:
            resp = requests.post(
                f"{OLLAMA_URL}/api/embed",
                json={"model": EMBED_MODEL, "input": texts},
                timeout=3600,
            )
            resp.raise_for_status()
            return np.array(resp.json()["embeddings"], dtype="float32")
        except Exception as e:
            print(f"  ! embed attempt {attempt + 1}/{retries} failed: {e}", flush=True)
            if attempt < retries - 1:
                time.sleep(5)
    raise


def embed_batch(texts, batch=EMBED_BATCH):
    parts = [embed(texts[i:i + batch]) for i in range(0, len(texts), batch)]
    return np.vstack(parts)


def save_index():
    INDEX_DIR.mkdir(parents=True, exist_ok=True)
    tmp = INDEX_DIR / "faiss.index.tmp"
    faiss.write_index(INDEX, str(tmp))
    os.replace(tmp, INDEX_DIR / "faiss.index")
    tmpj = INDEX_DIR / "meta.json.tmp"
    tmpj.write_text(json.dumps(META, ensure_ascii=False), encoding="utf-8")
    os.replace(tmpj, INDEX_DIR / "meta.json")


def load_index():
    global INDEX, META
    try:
        if (INDEX_DIR / "faiss.index").exists():
            INDEX = faiss.read_index(str(INDEX_DIR / "faiss.index"))
            META = json.loads((INDEX_DIR / "meta.json").read_text(encoding="utf-8"))
    except Exception as e:
        print("Не удалось прочитать индекс, пересоберу из кэша:", repr(e), flush=True)
        INDEX, META = None, []
    build_bm25()


def file_hash(p: Path) -> str:
    h = hashlib.md5()
    with open(p, "rb") as f:
        for block in iter(lambda: f.read(1 << 20), b""):
            h.update(block)
    return h.hexdigest()


def load_cache_entry(key: str):
    jp = CACHE_DIR / (key + ".json")
    vp = CACHE_DIR / (key + ".npy")
    if not jp.exists() or not vp.exists():
        return None
    try:
        meta = json.loads(jp.read_text(encoding="utf-8"))
        vecs = np.load(vp)
        return meta, vecs
    except Exception:
        return None


def save_cache_entry(key: str, meta: dict, vecs: np.ndarray):
    CACHE_DIR.mkdir(parents=True, exist_ok=True)
    (CACHE_DIR / (key + ".json")).write_text(
        json.dumps(meta, ensure_ascii=False), encoding="utf-8"
    )
    import io
    buf = io.BytesIO()
    np.save(buf, vecs)
    tmp = CACHE_DIR / (key + ".npy.tmp")
    tmp.write_bytes(buf.getvalue())
    os.replace(tmp, CACHE_DIR / (key + ".npy"))


@mcp.tool()
def reindex_documents() -> str:
    """Инкрементально перестраивает индекс."""
    with _reindex_lock:
        return _reindex_impl()


def _reindex_impl() -> str:
    global INDEX, META, LAST_STATE
    files = [p for p in DOCS_DIR.rglob("*") if p.is_file()]
    good = [p for p in files if p.suffix.lower() in SUPPORTED]
    skipped = sorted(
        {p.relative_to(DOCS_DIR).as_posix() for p in files if p.suffix.lower() not in SUPPORTED}
    )

    state = {}
    for p in sorted(good):
        state[p.relative_to(DOCS_DIR).as_posix()] = file_hash(p)
    if state == LAST_STATE and INDEX is not None:
        return "Изменений нет."

    new_meta, all_vecs = [], []
    reused = reembedded = 0
    seen_keys = set()
    total = len(state)
    start_time = time.time()
    processed = 0

    for idx, (rel, h) in enumerate(state.items(), 1):
        key = hashlib.sha1(rel.encode("utf-8")).hexdigest()
        seen_keys.add(key)
        entry = load_cache_entry(key)
        if entry and entry[0].get("hash") == h:
            meta, vecs = entry
            reused += 1
            print(f"[{idx}/{total}] {rel} — из кэша", flush=True)
        else:
            print(f"[{idx}/{total}] {rel} — читаю и считаю векторы...", flush=True)
            t0 = time.time()
            try:
                text = extract_text(DOCS_DIR / rel)
            except Exception as e:
                print(f"  ! ошибка чтения: {e}", flush=True)
                skipped.add(rel)
                continue
            chunks = chunk_text(text)
            if not chunks:
                continue
            vecs = embed_batch(chunks)
            meta = {"file": rel, "hash": h, "chunks": chunks}
            save_cache_entry(key, meta, vecs)
            reembedded += 1
            print(f"  -> {len(chunks)} кусков за {time.time() - t0:.1f}с", flush=True)
        for c in meta["chunks"]:
            new_meta.append({"file": rel, "text": c})
        all_vecs.append(vecs)
        processed += 1
        if processed > 1:
            avg = (time.time() - start_time) / processed
            eta = avg * (total - processed)
            print(f"  ETA: {int(eta // 60)}м {int(eta % 60)}с", flush=True)

    for jp in CACHE_DIR.glob("*.json"):
        if jp.stem not in seen_keys:
            npy = CACHE_DIR / (jp.stem + ".npy")
            jp.unlink(missing_ok=True)
            if npy.exists():
                npy.unlink()

    LAST_STATE = state
    if not new_meta:
        INDEX, META = None, []
        return "В папке " + str(DOCS_DIR) + " не найдено документов (.txt,.md,.pdf,.docx)."

    print("Собираю индекс...", flush=True)
    print(f"  vstack: {sum(len(v) for v in all_vecs)} векторов...", flush=True)
    X = np.vstack(all_vecs)
    print(f"  vstack done: {X.shape}", flush=True)
    print("  normalize...", flush=True)
    faiss.normalize_L2(X)
    print("  normalize done", flush=True)
    print(f"  создаю IndexFlatIP размерности {X.shape[1]}...", flush=True)
    INDEX = faiss.IndexFlatIP(X.shape[1])
    print("  add в индекс...", flush=True)
    INDEX.add(X)
    print(f"  add done, всего {INDEX.ntotal}", flush=True)
    META = new_meta
    print("  build_bm25...", flush=True)
    build_bm25()
    print("  save_index...", flush=True)
    save_index()
    print("  save done", flush=True)
    msg = (f"Готово. Файлов: {reused + reembedded}, кусков: {len(META)}. "
           f"Заново обработано: {reembedded}, переиспользовано: {reused}.")
    if skipped:
        msg += " Пропущены: " + ", ".join(skipped)
    return msg


@mcp.tool()
def search_documents(query: str, top_k: int = 20) -> str:
    """Гибридный поиск (смысл + ключевые слова)."""
    if INDEX is None or INDEX.ntotal == 0:
        return "Индекс пуст. Сначала вызовите reindex_documents."
    n = INDEX.ntotal
    q = embed([query])
    faiss.normalize_L2(q)
    D, I = INDEX.search(q, min(50, n))
    rrf = Counter()
    for rank, i in enumerate(I[0]):
        if i >= 0:
            rrf[int(i)] += 1.0 / (61 + rank)
    if BM25_OBJ is not None:
        bs = BM25_OBJ.score(query)
        order = sorted(range(n), key=lambda i: -bs[i])[:50]
        for rank, i in enumerate(order):
            if bs[i] > 0:
                rrf[i] += 1.0 / (61 + rank)
    out = []
    for i, sc in rrf.most_common(min(top_k, n)):
        out.append(f"[Файл: {META[i]['file']}] (score {sc:.3f})\n{META[i]['text']}")
    return "\n\n---\n\n".join(out)


@mcp.tool()
def list_documents() -> str:
    """Показывает, какие файлы проиндексированы."""
    if not META:
        return "Индекс пуст."
    c = Counter(m["file"] for m in META)
    return "\n".join(f"{name}: {n} кусков" for name, n in c.items())


class DocsHandler(FileSystemEventHandler):
    def __init__(self):
        self.timer = None

    def on_any_event(self, event):
        if event.is_directory:
            return
        if Path(event.src_path).suffix.lower() not in SUPPORTED:
            return
        if self.timer is not None:
            self.timer.cancel()
        self.timer = threading.Timer(REINDEX_DELAY, self._run)
        self.timer.start()

    def _run(self):
        try:
            print("Автообновление индекса...", flush=True)
            print(reindex_documents(), flush=True)
        except Exception as e:
            print("Ошибка автообновления:", e, flush=True)


def startup_catchup():
    try:
        print(reindex_documents(), flush=True)
    except Exception as e:
        print("Ошибка стартовой индексации:", e, flush=True)


if __name__ == "__main__":
    DOCS_DIR.mkdir(parents=True, exist_ok=True)
    CACHE_DIR.mkdir(parents=True, exist_ok=True)
    load_index()
    observer = Observer()
    observer.schedule(DocsHandler(), str(DOCS_DIR), recursive=True)
    observer.start()
    print(f"=== doc-search запущен на {HOST}:{PORT} ===", flush=True)
    print(f"=== Папка документов: {DOCS_DIR} ===", flush=True)
    print(f"=== Embedding: {EMBED_MODEL} | OCR: {OCR_ENABLED} ===", flush=True)
    threading.Thread(target=startup_catchup, daemon=True).start()
    try:
        mcp.run(transport="streamable-http")
    finally:
        observer.stop()
    observer.join()
